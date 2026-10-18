"""Core conversion engine. Pure Python, no UI, works on macOS and Linux.

Tools used (all free): Pillow, pillow-heif, PyMuPDF, ffmpeg, LibreOffice.
"""
from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

IMAGE = {"heic", "heif", "png", "jpg", "jpeg", "webp", "gif", "tiff", "tif", "bmp"}
VIDEO = {"mov", "mp4", "mkv", "avi", "webm", "m4v"}
AUDIO = {"wav", "mp3", "m4a", "aac", "flac", "ogg", "aiff", "aif"}
OFFICE = {"docx", "doc", "xlsx", "xls", "odt", "rtf", "txt"}
SLIDES = {"pptx", "ppt"}
PDF = {"pdf"}

TARGETS = {
    "image": ["png", "jpg", "webp", "pdf", "gif", "pptx"],
    "video": ["mp4", "mov", "webm", "gif", "mp3"],
    "audio": ["mp3", "m4a", "wav", "flac", "ogg"],
    "office": ["pdf"],
    "pdf": ["png", "txt", "pptx"],
    "slides": ["pdf", "png", "jpg", "pptx", "ppt"],
}


class ConversionError(Exception):
    pass


def kind_of(path: Path) -> str | None:
    ext = path.suffix.lower().lstrip(".")
    if ext in IMAGE:
        return "image"
    if ext in VIDEO:
        return "video"
    if ext in AUDIO:
        return "audio"
    if ext in SLIDES:
        return "slides"
    if ext in OFFICE:
        return "office"
    if ext in PDF:
        return "pdf"
    return None


def targets_for(path: Path) -> list[str]:
    """Formats offered for this file, never the format it already is."""
    kind = kind_of(path)
    if kind is None:
        return []
    own = path.suffix.lower().lstrip(".")
    own = "jpg" if own == "jpeg" else own
    return [t for t in TARGETS[kind] if t != own]


def unique_path(src: Path, ext: str) -> Path:
    """Output beside the original. Never overwrites an existing file."""
    out = src.with_suffix(f".{ext}")
    n = 1
    while out.exists():
        out = src.with_name(f"{src.stem} ({n}).{ext}")
        n += 1
    return out


def _run(cmd: list[str]) -> None:
    try:
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=3600)
    except FileNotFoundError:
        raise ConversionError(f"'{cmd[0]}' is not installed. See README.")
    if r.returncode != 0:
        raise ConversionError(r.stderr.strip().splitlines()[-1] if r.stderr.strip() else "failed")


def _image(src: Path, ext: str) -> list[Path]:
    from PIL import Image

    try:
        import pillow_heif

        pillow_heif.register_heif_opener()
    except ImportError:
        pass
    out = unique_path(src, ext)
    if ext == "pptx":
        with tempfile.TemporaryDirectory() as tmp, Image.open(src) as im:
            png = Path(tmp) / "slide.png"
            im.convert("RGB").save(png)
            _build_pptx([png], out)
        return [out]
    with Image.open(src) as im:
        if ext in ("jpg", "pdf") and im.mode in ("RGBA", "P", "LA"):
            im = im.convert("RGB")
        fmt = {"jpg": "JPEG", "png": "PNG", "webp": "WEBP", "gif": "GIF", "pdf": "PDF"}[ext]
        im.save(out, fmt)
    return [out]


def _ffmpeg(src: Path, ext: str) -> list[Path]:
    out = unique_path(src, ext)
    cmd = ["ffmpeg", "-y", "-loglevel", "error", "-i", str(src)]
    if ext == "gif":
        cmd += ["-vf", "fps=12,scale=720:-1:flags=lanczos", "-loop", "0"]
    elif ext == "mp3":
        cmd += ["-vn", "-q:a", "2"]
    elif ext in ("m4a",):
        cmd += ["-vn", "-c:a", "aac", "-b:a", "192k"]
    elif ext == "mp4":
        cmd += ["-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", "-movflags", "+faststart"]
    elif ext == "mov":
        cmd += ["-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac"]
    elif ext == "webm":
        cmd += ["-c:v", "libvpx-vp9", "-c:a", "libopus"]
    elif ext == "flac":
        cmd += ["-vn", "-c:a", "flac"]
    elif ext == "ogg":
        cmd += ["-vn", "-c:a", "libvorbis", "-q:a", "5"]
    elif ext == "wav":
        cmd += ["-vn"]
    cmd.append(str(out))
    _run(cmd)
    return [out]


def _soffice() -> str:
    for name in ("soffice", "libreoffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if shutil.which(name) or Path(name).exists():
            return name
    raise ConversionError("LibreOffice is not installed (brew install --cask libreoffice).")


def _office(src: Path, ext: str) -> list[Path]:
    out = unique_path(src, ext)
    with tempfile.TemporaryDirectory() as tmp:
        _run([_soffice(), "--headless", "--convert-to", ext, "--outdir", tmp, str(src)])
        made = Path(tmp) / f"{src.stem}.{ext}"
        if not made.exists():
            raise ConversionError("LibreOffice produced no file.")
        shutil.move(str(made), out)
    return [out]


def _build_pptx(images: list[Path], out: Path) -> None:
    """One full-bleed slide per image; slide shape follows the first image."""
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()
    prs.slide_width = Emu(12192000)  # 13.333 in
    with Image.open(images[0]) as first:
        prs.slide_height = Emu(int(prs.slide_width * first.height / first.width))
    blank = prs.slide_layouts[6]
    for img in images:
        with Image.open(img) as im:
            w, h = im.size
        scale = min(prs.slide_width / w, prs.slide_height / h)
        pw, ph = int(w * scale), int(h * scale)
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(img), (prs.slide_width - pw) // 2, (prs.slide_height - ph) // 2, pw, ph
        )
    prs.save(out)


def _render_pdf_pages(pdf: Path, folder: Path, ext: str, dpi: int) -> list[Path]:
    import pymupdf

    pages: list[Path] = []
    with pymupdf.open(pdf) as doc:
        for i, page in enumerate(doc, start=1):
            dest = folder / f"p{i:04d}.{ext}"
            page.get_pixmap(dpi=dpi, alpha=False).save(dest)
            pages.append(dest)
    return pages


def _slides(src: Path, ext: str) -> list[Path]:
    """PowerPoint input: PDF, pptx/ppt via LibreOffice, png/jpg one file per slide."""
    if ext in ("pdf", "pptx", "ppt"):
        return _office(src, ext)
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        _run([_soffice(), "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(src)])
        pdf = tmp / f"{src.stem}.pdf"
        if not pdf.exists():
            raise ConversionError("LibreOffice produced no file.")
        pages = _render_pdf_pages(pdf, tmp, ext, dpi=150)
        outs = []
        for i, page in enumerate(pages, start=1):
            base = src if len(pages) == 1 else src.with_name(f"{src.stem}-p{i}{src.suffix}")
            out = unique_path(base, ext)
            shutil.move(str(page), out)
            outs.append(out)
    return outs


def _pdf(src: Path, ext: str) -> list[Path]:
    import pymupdf

    doc = pymupdf.open(src)
    outs: list[Path] = []
    if ext == "txt":
        out = unique_path(src, "txt")
        out.write_text("\n\n".join(p.get_text() for p in doc), encoding="utf-8")
        return [out]
    if ext == "pptx":
        out = unique_path(src, "pptx")
        with tempfile.TemporaryDirectory() as tmp:
            _build_pptx(_render_pdf_pages(src, Path(tmp), "png", dpi=150), out)
        return [out]
    # png: one file per page
    for i, page in enumerate(doc, start=1):
        base = src if len(doc) == 1 else src.with_name(f"{src.stem}-p{i}.pdf")
        out = unique_path(base, "png")
        page.get_pixmap(dpi=200).save(out)
        outs.append(out)
    return outs


def convert(src: str | Path, ext: str) -> list[Path]:
    """Convert `src` to `ext`. Returns the created files. Raises ConversionError."""
    src = Path(src)
    ext = ext.lower().lstrip(".")
    if not src.is_file():
        raise ConversionError(f"Not a file: {src.name}")
    if ext not in targets_for(src):
        raise ConversionError(f"Can't convert {src.suffix} to .{ext}")
    kind = kind_of(src)
    try:
        if kind == "image":
            return _image(src, ext)
        if kind in ("video", "audio"):
            return _ffmpeg(src, ext)
        if kind == "slides":
            return _slides(src, ext)
        if kind == "office":
            return _office(src, ext)
        return _pdf(src, ext)
    except ConversionError:
        raise
    except Exception as e:  # surface library errors as one clean type
        raise ConversionError(str(e)) from e
